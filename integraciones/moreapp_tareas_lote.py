"""Carga masiva de tareas MoreApp (Excel/CSV y clientes de la intranet)."""

from __future__ import annotations

import csv
import io
import re
from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional, Tuple

from .moreapp_fuentes import fold_header

MAX_LOTE = 40
MAX_CANTIDAD = 20
MAX_FILAS_LISTADO = 15

_COL_ALIASES = {
    "formulario": ("formulario", "form", "formid", "idformulario", "nombreformulario"),
    "destinatario": (
        "destinatario",
        "destinatarios",
        "correo",
        "email",
        "usuario",
        "grupo",
        "tecnico",
        "técnico",
    ),
    "mensaje": ("mensaje", "message", "nota", "descripcion", "descripción"),
    "fecha": ("fecha", "fechalimite", "vence", "vencimiento", "due", "duedate"),
    "ubicacion": ("ubicacion", "ubicación", "location", "direccion", "dirección"),
    "cliente": ("cliente", "cuenta", "numerocliente", "nrocliente", "numero"),
}

_LATLNG = re.compile(
    r"^\s*(-?\d+(?:\.\d+)?)\s*[,;\s]\s*(-?\d+(?:\.\d+)?)\s*$"
)


def _norm(value: Any) -> str:
    return fold_header(str(value or ""))


def _map_headers(headers: Iterable[Any]) -> Dict[int, str]:
    mapping: Dict[int, str] = {}
    seen = set()
    for idx, raw in enumerate(headers):
        folded = _norm(raw)
        for canonical, aliases in _COL_ALIASES.items():
            if folded in aliases and canonical not in seen:
                mapping[idx] = canonical
                seen.add(canonical)
                break
    return mapping


def _row_dict(values: List[Any], mapping: Dict[int, str]) -> Dict[str, str]:
    row = {key: "" for key in _COL_ALIASES}
    for idx, canonical in mapping.items():
        if idx < len(values):
            row[canonical] = str(values[idx] or "").strip()
    return row


def _filas_desde_headers(
    headers: Iterable[Any], rows_iter: Iterable[Any]
) -> Optional[Tuple[List[Dict[str, str]], int]]:
    mapping = _map_headers(headers or ())
    if "formulario" not in mapping.values() or "destinatario" not in mapping.values():
        return None
    out: List[Dict[str, str]] = []
    total = 0
    for values in rows_iter:
        cells = list(values or ())
        if not any(str(v or "").strip() for v in cells):
            continue
        total += 1
        if len(out) < MAX_LOTE:
            out.append(_row_dict(cells, mapping))
    return out, total


def parse_lote_csv(content: bytes) -> Tuple[List[Dict[str, str]], int]:
    text = ""
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text.strip():
        return [], 0
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows_iter = list(reader)
    if not rows_iter:
        return [], 0
    parsed = _filas_desde_headers(rows_iter[0], rows_iter[1:])
    if parsed is None:
        raise ValueError("El archivo debe incluir las columnas formulario y destinatario.")
    return parsed


def parse_lote_xlsx(content: bytes) -> Tuple[List[Dict[str, str]], int]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("No se pudo leer Excel. Usa CSV o instala openpyxl.") from exc
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheetnames = list(wb.sheetnames) if hasattr(wb, "sheetnames") else []
    ordered: List[str] = []
    if "Tareas" in sheetnames:
        ordered.append("Tareas")
    active_title = str(getattr(wb.active, "title", "") or "")
    if active_title and active_title not in ordered:
        ordered.append(active_title)
    for name in sheetnames:
        if name not in ordered:
            ordered.append(name)
    for name in ordered:
        hoja = wb[name]
        rows_iter = hoja.iter_rows(values_only=True)
        try:
            headers = next(rows_iter)
        except StopIteration:
            continue
        parsed = _filas_desde_headers(headers or (), rows_iter)
        if parsed is None:
            continue
        return parsed
    raise ValueError("El archivo debe incluir las columnas formulario y destinatario.")


def parse_lote_archivo(nombre: str, content: bytes) -> Tuple[List[Dict[str, str]], int]:
    lower = (nombre or "").lower()
    if lower.endswith((".xlsx", ".xlsm")):
        rows, total = parse_lote_xlsx(content)
    else:
        rows, total = parse_lote_csv(content)
    return rows, total


def parse_fecha_lote(value: str) -> str:
    raw = (value or "").strip()
    if not raw:
        return ""
    for fmt in ("%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            parsed = datetime.strptime(raw, fmt)
            return parsed.strftime("%Y-%m-%dT%H:%M")
        except ValueError:
            continue
    return ""


def parse_ubicacion_lote(value: str) -> Tuple[str, str, str]:
    raw = (value or "").strip()
    if not raw:
        return "", "", ""
    match = _LATLNG.match(raw)
    if match:
        return raw, match.group(1), match.group(2)
    return raw, "", ""


def split_destinatarios(value: str) -> List[str]:
    parts = re.split(r"[;,]", value or "")
    return [item.strip() for item in parts if item.strip()]


def cuenta_de_fila(row: Dict[str, Any]) -> str:
    if not isinstance(row, dict):
        return ""
    for key, val in row.items():
        folded = _norm(key)
        if folded in ("cuenta", "numerocliente", "nrocliente", "cliente", "numero"):
            return str(val or "").strip()
    return str(row.get("CUENTA") or row.get("cuenta") or "").strip()


def fila_cliente_en_fuente(
    filas: Iterable[Dict[str, Any]], numero_cliente: str
) -> Optional[Dict[str, Any]]:
    wanted = _norm(numero_cliente)
    if not wanted:
        return None
    for row in filas:
        if not isinstance(row, dict):
            continue
        if _norm(cuenta_de_fila(row)) == wanted:
            return row
    return None


_PLANTILLA_COLUMNAS = (
    "formulario",
    "destinatario",
    "mensaje",
    "fecha",
    "ubicacion",
    "cliente",
)


def plantilla_csv() -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(_PLANTILLA_COLUMNAS)
    return buffer.getvalue().encode("utf-8-sig")


def plantilla_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.comments import Comment

    ayudas = {
        "formulario": "Nombre o ID del formulario publicado en MoreApp.",
        "destinatario": "Correo de un usuario MoreApp, o el nombre del grupo.",
        "mensaje": "Texto corto que ve quien recibe la tarea.",
        "fecha": "Fecha límite opcional, por ejemplo 2026-10-23 09:00.",
        "ubicacion": "Dirección, o latitud y longitud separadas por coma.",
        "cliente": "Número de cliente (CUENTA). Si coincide el Search, se precarga.",
    }
    anchos = {
        "formulario": 28,
        "destinatario": 32,
        "mensaje": 36,
        "fecha": 22,
        "ubicacion": 28,
        "cliente": 18,
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "Tareas"
    header_fill = PatternFill("solid", fgColor="1A3A5C")
    header_font = Font(bold=True, color="FFFFFF")
    for col, nombre in enumerate(_PLANTILLA_COLUMNAS, start=1):
        cell = ws.cell(row=1, column=col, value=nombre)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.comment = Comment(ayudas[nombre], "DELCO")
        ws.column_dimensions[get_column_letter(col)].width = anchos[nombre]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:F1"

    guia = wb.create_sheet("Instrucciones")
    guia["A1"] = "Columna"
    guia["B1"] = "Qué completar"
    guia["A1"].font = header_font
    guia["B1"].font = header_font
    guia["A1"].fill = header_fill
    guia["B1"].fill = header_fill
    for idx, nombre in enumerate(_PLANTILLA_COLUMNAS, start=2):
        guia.cell(row=idx, column=1, value=nombre)
        guia.cell(row=idx, column=2, value=ayudas[nombre])
    guia.column_dimensions["A"].width = 18
    guia.column_dimensions["B"].width = 70
    guia["A8"] = "Una fila = una tarea. Completa desde la fila 2. No juntes varios campos en la misma celda."
    guia.merge_cells("A8:B8")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def resolver_formulario(
    valor: str, formularios: List[Dict[str, Any]]
) -> Optional[Dict[str, Any]]:
    raw = (valor or "").strip()
    if not raw:
        return None
    folded = _norm(raw)
    homonimos: List[Dict[str, Any]] = []
    for form in formularios:
        fid = str(form.get("id") or "").strip()
        name = str(form.get("name") or "").strip()
        if raw == fid:
            return form
        if name and _norm(name) == folded:
            homonimos.append(form)
    if len(homonimos) == 1:
        return homonimos[0]
    return None


def resolver_destinatario(
    token: str,
    usuarios: List[Dict[str, Any]],
    grupos: List[Dict[str, Any]],
) -> Tuple[List[str], List[str], str]:
    raw = (token or "").strip()
    if not raw:
        return [], [], "Falta el destinatario."
    folded = _norm(raw)
    if "@" in raw:
        wanted = raw.lower()
        for user in usuarios:
            email = str(user.get("email") or "").strip()
            if email.lower() == wanted:
                return [email], [], ""
        return [], [], "El correo %s no está en MoreApp." % raw
    for group in grupos:
        gid = str(group.get("id") or "").strip()
        nombre = str(group.get("nombre") or "").strip()
        nombre_base = nombre.split(" · ", 1)[0].strip()
        aliases = [item for item in (nombre, nombre_base) if item]
        if raw == gid or any(_norm(item) == folded for item in aliases):
            if gid:
                return [], [gid], ""
    for user in usuarios:
        email = str(user.get("email") or "").strip()
        nombre = str(user.get("nombre") or "").strip()
        username = str(user.get("username") or "").strip()
        if email and (_norm(nombre) == folded or _norm(username) == folded):
            return [email], [], ""
    return [], [], "No se reconoció el destinatario %s." % raw


def unicos(items: Iterable[str]) -> List[str]:
    out: List[str] = []
    seen = set()
    for item in items:
        value = str(item or "").strip()
        if value and value not in seen:
            seen.add(value)
            out.append(value)
    return out


def destino_listado(token: str) -> Tuple[str, str]:
    raw = (token or "").strip()
    if raw.startswith("user:"):
        return "user", raw[5:].strip()
    if raw.startswith("group:"):
        return "group", raw[6:].strip()
    return "", raw


def parse_filas_listado(post) -> Tuple[List[Dict[str, Any]], str]:
    forms = list(post.getlist("form_id"))
    dests = list(post.getlist("destinatario"))
    cants = list(post.getlist("cantidad"))
    n = max(len(forms), len(dests), len(cants))
    if n > MAX_FILAS_LISTADO:
        return [], "Máximo %s filas en el listado." % MAX_FILAS_LISTADO
    filas: List[Dict[str, Any]] = []
    total = 0
    for idx in range(n):
        form_id = str(forms[idx] if idx < len(forms) else "").strip()
        dest = str(dests[idx] if idx < len(dests) else "").strip()
        raw_cant = str(cants[idx] if idx < len(cants) else "").strip()
        if not form_id and not dest:
            continue
        if not form_id or not dest:
            return [], "Cada fila necesita formulario y persona o grupo."
        try:
            cantidad = int(raw_cant or "0")
        except (TypeError, ValueError):
            return [], "La cantidad debe ser un número entero."
        if cantidad < 1 or cantidad > MAX_CANTIDAD:
            return [], "La cantidad por fila va de 1 a %s." % MAX_CANTIDAD
        total += cantidad
        if total > MAX_LOTE:
            return [], "El listado no puede superar %s tareas." % MAX_LOTE
        filas.append({"form_id": form_id, "destinatario": dest, "cantidad": cantidad})
    if not filas:
        return [], "Agrega al menos una fila con formulario, persona o grupo y cantidad."
    return filas, ""
