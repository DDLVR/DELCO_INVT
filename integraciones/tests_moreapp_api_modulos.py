# coding: utf-8
import json
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from unittest import mock

Usuario = get_user_model()


class MoreAppApiModulosTests(TestCase):
    def setUp(self):
        self.password = 'Seguro1234'
        self.admin = Usuario.objects.create_user(
            rut='33333333-3',
            email='admin-api@delco.cl',
            password=self.password,
            nombre='Admin',
            apellido='Api',
            nombre_interno='admin_api',
            rol='ADMIN',
            is_active=True,
        )
        self.tecnico = Usuario.objects.create_user(
            rut='44444444-4',
            email='tec-api@delco.cl',
            password=self.password,
            nombre='Tec',
            apellido='Api',
            nombre_interno='tec_api',
            rol='TECNICO',
            is_active=True,
        )

    def test_catalogo_requiere_login(self):
        resp = self.client.get(reverse('moreapp_api_catalogo'))
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/login', resp.url)

    def test_tecnico_no_accede_catalogo(self):
        self.client.login(username=self.tecnico.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_api_catalogo'))
        self.assertEqual(resp.status_code, 403)

    def test_administrativo_no_accede_catalogo(self):
        adminis = Usuario.objects.create_user(
            rut='66666666-6',
            email='adm-api@delco.cl',
            password=self.password,
            nombre='Adm',
            apellido='Api',
            nombre_interno='adm_api',
            rol='ADMINISTRATIVO',
            is_active=True,
        )
        self.client.login(username=adminis.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_api_catalogo'))
        self.assertEqual(resp.status_code, 403)
        dash = self.client.get(reverse('dashboard'))
        body = dash.content.decode('utf-8')
        self.assertNotIn('Catálogo MoreApp', body)
        self.assertIn('Formularios asignados', body)

    def test_admin_ve_catalogo_en_administracion(self):
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('dashboard'))
        body = resp.content.decode('utf-8')
        moreapp_block = body.split('MoreApp y reportes', 1)[1].split('Control', 1)[0]
        admin_block = body.split('Administración', 1)[1]
        self.assertNotIn('Catálogo MoreApp', moreapp_block)
        self.assertIn('Formularios asignados', moreapp_block)
        self.assertIn('Catálogo MoreApp', admin_block)
        operacion = body.split('Operación', 1)[1].split('MoreApp y reportes', 1)[0]
        self.assertIn('Tareas MoreApp', operacion)
        self.assertIn(reverse('moreapp_tareas_list'), operacion)

    @override_settings(MOREAPP_API_ENABLED=False)
    def test_json_forms_sin_api_enabled(self):
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_v1_visualizacion_forms_list'))
        self.assertEqual(resp.status_code, 400)
        payload = resp.json()
        self.assertFalse(payload['success'])
        self.assertEqual(payload['module'], 'visualizacion')

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_visualizacion_views.list_forms')
    def test_json_forms_ok_no_expone_api_key(self, mocked_list):
        mocked_list.return_value = [{'id': 'abc123abc123abc123abc123', 'name': 'OT Terreno'}]
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_v1_visualizacion_forms_list'))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertNotIn('test-key', body)
        payload = resp.json()
        self.assertTrue(payload['success'])
        self.assertEqual(payload['module'], 'visualizacion')

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_datos_views.list_users')
    def test_json_users_modulo_datos(self, mocked_users):
        mocked_users.return_value = [{'id': 'u1', 'emailAddress': 'tec@delco.cl'}]
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_v1_datos_users_list'))
        self.assertEqual(resp.status_code, 200)
        payload = resp.json()
        self.assertEqual(payload['module'], 'datos')
        self.assertEqual(payload['operation'], 'users.list')

    def test_form_id_invalido_no_llama_moreapp(self):
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(
            reverse('moreapp_v1_visualizacion_form_detail', args=['invalido!formulario'])
        )
        self.assertEqual(resp.status_code, 400)

    def test_registros_requiere_login(self):
        resp = self.client.get(
            reverse('moreapp_api_form_registros', args=['abc123abc123abc123abc123'])
        )
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/login', resp.url)

    def test_tecnico_no_accede_registros(self):
        self.client.login(username=self.tecnico.rut, password=self.password)
        resp = self.client.get(
            reverse('moreapp_api_form_registros', args=['abc123abc123abc123abc123'])
        )
        self.assertEqual(resp.status_code, 403)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_registros_views.get_form_version')
    @mock.patch('integraciones.moreapp_registros_views.list_submissions')
    @mock.patch('integraciones.moreapp_registros_views.get_form')
    def test_registros_lista_aplana_info(self, mocked_form, mocked_subs, mocked_version):
        mocked_form.return_value = {
            'id': 'abc123abc123abc123abc123',
            'name': 'Reposición de Suministro',
            'status': 'ACTIVE',
            'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
        }
        mocked_version.return_value = {
            'id': 'ver123ver123ver123ver123',
            'fields': [
                {
                    'widget': 'com.moreapps:number:1',
                    'properties': {'label_text': 'Num Cliente', 'data_name': 'numCliente'},
                }
            ],
        }
        mocked_subs.return_value = [
            {
                'id': 'sub123sub123sub123sub123',
                'info': {'userName': 'tec@delco.cl', 'date': '2026-10-14T12:00:00Z'},
                'meta': {'lastUpdated': '2026-10-14T12:01:00Z'},
                'data': {
                    'numCliente': '9852366',
                    'tecnicoResponsable': {'NOMBRES': 'Pedro Jalvear'},
                },
            }
        ]
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(
            reverse('moreapp_api_form_registros', args=['abc123abc123abc123abc123'])
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertNotIn('test-key', body)
        self.assertIn('9852366', body)
        self.assertIn('Pedro Jalvear', body)
        self.assertIn('Técnico responsable', body)
        self.assertIn('N° cliente', body)
        self.assertIn('14/10/2026', body)
        self.assertIn('Ver contenido', body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_registros_views.get_form_version')
    @mock.patch('integraciones.moreapp_registros_views.get_submission')
    @mock.patch('integraciones.moreapp_registros_views.get_form')
    def test_registro_detalle_muestra_campos(self, mocked_form, mocked_sub, mocked_version):
        mocked_form.return_value = {
            'id': 'abc123abc123abc123abc123',
            'name': 'Reposición de Suministro',
            'status': 'ACTIVE',
            'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
        }
        mocked_version.return_value = {
            'fields': [
                {
                    'widget': 'com.moreapps:number:1',
                    'properties': {'label_text': 'Num Cliente', 'data_name': 'numCliente'},
                }
            ]
        }
        mocked_sub.return_value = {
            'id': 'sub123sub123sub123sub123',
            'info': {'userName': 'tec@delco.cl'},
            'meta': {'created': '2026-10-14T12:00:00Z'},
            'data': {
                'numCliente': '9852366',
                'nota': 'Medidor ok',
                'fachada': 'gridfs://registrationFiles/7f555c8d-6a67-439c-ad41-da32206657b9',
            },
        }
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(
            reverse(
                'moreapp_api_registro',
                args=['abc123abc123abc123abc123', 'sub123sub123sub123sub123'],
            )
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('Num Cliente', body)
        self.assertIn('9852366', body)
        self.assertIn('Medidor ok', body)
        self.assertIn('7f555c8d-6a67-439c-ad41-da32206657b9', body)
        self.assertNotIn('test-key', body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_registros_views.get_latest_submission')
    def test_ultimo_registro_devuelve_fecha(self, mocked_latest):
        mocked_latest.return_value = {
            'totalSize': 1,
            'elements': [
                {
                    'id': 'sub123sub123sub123sub123',
                    'info': {'date': 1739452535990, 'userId': 'tec@delco.cl'},
                }
            ],
        }
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(
            reverse('moreapp_api_form_ultimo_registro', args=['abc123abc123abc123abc123'])
        )
        self.assertEqual(resp.status_code, 200)
        payload = resp.json()
        self.assertTrue(payload['ok'])
        self.assertEqual(payload['date'], '13/02/2025 10:15')
        self.assertEqual(payload['day'], '2025-02-13')
        self.assertNotIn('test-key', resp.content.decode('utf-8'))

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_catalogo_views.list_all_forms')
    def test_catalogo_filtra_por_nombre(self, mocked_forms):
        mocked_forms.return_value = [
            {'id': 'abc123abc123abc123abc123', 'meta': {'name': 'Reposición de Suministro'}, 'status': 'ACTIVE'},
            {'id': 'def456def456def456def456', 'meta': {'name': 'Mantención medidor'}, 'status': 'ACTIVE'},
        ]
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_api_catalogo'), {'query': 'reposicion'})
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('Reposición de Suministro', body)
        self.assertIn('Mantención medidor', body)
        self.assertIn('js-catalogo-buscar', body)
        self.assertIn('value="reposicion"', body)
        self.assertIn('data-search=', body)
        self.assertIn('name="fecha_desde"', body)
        self.assertIn('name="fecha_hasta"', body)
        self.assertIn('name="estado"', body)
        self.assertIn('name="asignacion"', body)
        self.assertIn('Con asignados', body)
        self.assertIn('7 días', body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_catalogo_views.list_all_forms')
    def test_catalogo_pasa_fechas_a_registros(self, mocked_forms):
        mocked_forms.return_value = [
            {'id': 'abc123abc123abc123abc123', 'meta': {'name': 'Reposición de Suministro'}, 'status': 'ACTIVE'},
        ]
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(
            reverse('moreapp_api_catalogo'),
            {'fecha_desde': '2025-02-01', 'fecha_hasta': '2025-02-28'},
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('fecha_desde=2025-02-01', body)
        self.assertIn('fecha_hasta=2025-02-28', body)
        self.assertIn('name="fecha_desde"', body)
        self.assertIn('js-form-card', body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_catalogo_views.list_all_forms')
    def test_catalogo_ocultar_es_por_usuario(self, mocked_forms):
        from integraciones.models import MoreAppFormularioOculto

        form_id = 'abc123abc123abc123abc123'
        mocked_forms.return_value = [
            {'id': form_id, 'meta': {'name': 'Reposición de Suministro'}, 'status': 'ACTIVE'},
        ]
        otro = Usuario.objects.create_user(
            rut='55555555-5',
            email='admin2-api@delco.cl',
            password=self.password,
            nombre='Admin',
            apellido='Dos',
            nombre_interno='admin_api_2',
            rol='ADMIN',
            is_active=True,
        )
        self.client.login(username=self.admin.rut, password=self.password)
        hide = self.client.post(
            reverse('moreapp_api_form_ocultar', args=[form_id]),
            {'nombre': 'Reposición de Suministro'},
        )
        self.assertEqual(hide.status_code, 302)
        self.assertTrue(
            MoreAppFormularioOculto.objects.filter(usuario=self.admin, form_id=form_id).exists()
        )
        resp = self.client.get(reverse('moreapp_api_catalogo'))
        self.assertNotIn('Reposición de Suministro', resp.content.decode('utf-8'))
        shown = self.client.get(reverse('moreapp_api_catalogo'), {'mostrar_ocultos': '1'})
        self.assertIn('Reposición de Suministro', shown.content.decode('utf-8'))
        self.assertIn('Oculto en tu catálogo', shown.content.decode('utf-8'))
        self.client.logout()
        self.client.login(username=otro.rut, password=self.password)
        other = self.client.get(reverse('moreapp_api_catalogo'))
        self.assertIn('Reposición de Suministro', other.content.decode('utf-8'))
        restore = self.client.post(reverse('moreapp_api_form_mostrar', args=[form_id]))
        self.assertEqual(restore.status_code, 302)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_registros_views.get_form_version')
    @mock.patch('integraciones.moreapp_registros_views.list_submissions')
    @mock.patch('integraciones.moreapp_registros_views.get_form')
    def test_registros_filtra_por_fecha_y_nombre(self, mocked_form, mocked_subs, mocked_version):
        mocked_form.return_value = {
            'id': 'abc123abc123abc123abc123',
            'name': 'Reposición de Suministro',
            'status': 'ACTIVE',
            'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
        }
        mocked_version.return_value = {'fields': []}
        mocked_subs.return_value = {
            'totalSize': 1,
            'elements': [
                {
                    'id': 'sub123sub123sub123sub123',
                    'info': {'userName': 'Juan Perez', 'userId': 'juan@delco.cl', 'date': '2026-10-14T12:00:00Z'},
                    'data': {
                        'numCliente': 858678,
                        'tecnicoResponsable': {'NOMBRES': 'Juan Perez'},
                    },
                }
            ],
        }
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(
            reverse('moreapp_api_form_registros', args=['abc123abc123abc123abc123']),
            {
                'fecha_desde': '2026-10-01',
                'fecha_hasta': '2026-10-31',
                'tecnico': 'Juan',
                'num_cliente': '858678',
            },
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('Juan Perez', body)
        self.assertIn('858678', body)
        payload = mocked_subs.call_args.kwargs.get('payload') or mocked_subs.call_args[1].get('payload')
        query = payload['query']
        self.assertEqual(query[0]['path'], 'info.date')
        self.assertEqual(query[0]['type'], 'date')
        self.assertIn('start', query[0]['value'])
        self.assertIn('end', query[0]['value'])
        self.assertEqual(query[1]['path'], 'data.numCliente')
        self.assertEqual(query[1]['value'], 858678)
        self.assertIn('tecnico=Juan', body)
        self.assertIn('num_cliente=858678', body)
        self.assertIn('export=xlsx', body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_api.list_forms')
    def test_catalogo_junta_varias_paginas(self, mocked_forms):
        mocked_forms.side_effect = [
            [{'id': 'aaa123abc123abc123abc123', 'meta': {'name': 'Formulario A'}, 'status': 'ACTIVE'}],
            [{'id': 'bbb456def456def456def456', 'meta': {'name': 'Formulario B'}, 'status': 'ACTIVE'}],
            [],
        ]
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_api_catalogo'))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('Formulario A', body)
        self.assertIn('Formulario B', body)
        self.assertGreaterEqual(mocked_forms.call_count, 2)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_catalogo_views.get_form_version')
    @mock.patch('integraciones.moreapp_catalogo_views.list_form_versions')
    @mock.patch('integraciones.moreapp_catalogo_views.get_form')
    def test_formulario_version_fecha_ddmmaaaa(self, mocked_form, mocked_versions, mocked_version):
        mocked_form.return_value = {
            'id': 'abc123abc123abc123abc123',
            'name': 'Reposición de Suministro',
            'status': 'ACTIVE',
            'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
        }
        mocked_versions.return_value = [
            {
                'id': 'ver123ver123ver123ver123',
                'status': 'PUBLISHED',
                'meta': {'lastUpdated': '2026-10-14T12:00:00Z', 'lastUpdatedBy': 'editor'},
            }
        ]
        mocked_version.return_value = {'fields': []}
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_api_formulario', args=['abc123abc123abc123abc123']))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('14/10/2026', body)
        self.assertNotIn('2026-10-14T12:00:00Z', body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_catalogo_views.list_datasources')
    def test_datasources_tabla_columnas_y_fecha(self, mocked_ds):
        mocked_ds.return_value = [
            {
                'id': '6491e945a2abd076df165abc',
                'name': 'comunas',
                'searchType': 'LOCAL_FALLBACK',
                'enabled': True,
                'lastUpdated': 1687290265881,
                'columnMapping': [{'id': 'id', 'label': 'Id'}, {'id': 'comuna', 'label': 'Comuna'}],
            }
        ]
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_api_datasources'))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('comunas', body)
        self.assertIn('Comuna', body)
        self.assertIn('Última actualización', body)
        self.assertIn('20/06/2023', body)
        self.assertNotIn('1687290265881', body)
        self.assertIn('js-ds-buscar', body)
        self.assertIn('js-ds-card', body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_catalogo_views.list_datasource_entry_rows')
    @mock.patch('integraciones.moreapp_catalogo_views.get_datasource')
    def test_datasource_muestra_contenido_de_filas(self, mocked_get, mocked_rows):
        mocked_get.return_value = {
            'id': '6491e945a2abd076df165abc',
            'name': 'comunas',
            'columnMapping': [{'id': 'id', 'label': 'Id'}, {'id': 'comuna', 'label': 'Comuna'}],
        }
        mocked_rows.return_value = {
            'total': 2,
            'columns': [{'id': 'id', 'label': 'Id'}, {'id': 'comuna', 'label': 'Comuna'}],
            'rows': [
                {'id': '1', 'comuna': 'CONCHALI'},
                {'id': '2', 'comuna': 'SANTIAGO'},
            ],
            'fetched': 2,
            'truncated': False,
            'limit': 250,
        }
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_api_datasource', args=['6491e945a2abd076df165abc']))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('Contenido de la tabla', body)
        self.assertIn('CONCHALI', body)
        self.assertIn('SANTIAGO', body)
        self.assertIn('js-ds-filas-buscar', body)
        filtrado = self.client.get(
            reverse('moreapp_api_datasource', args=['6491e945a2abd076df165abc']),
            {'query': 'santiago'},
        )
        filtrado_body = filtrado.content.decode('utf-8')
        self.assertIn('SANTIAGO', filtrado_body)
        self.assertIn('CONCHALI', filtrado_body)
        self.assertIn('value="santiago"', filtrado_body)
        self.assertIn('js-ds-fila', filtrado_body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_catalogo_views.list_users')
    def test_usuarios_moreapp_filtra_correo(self, mocked_users):
        mocked_users.return_value = [
            {
                'id': 'u1',
                'emailAddress': 'juan@delco.cl',
                'username': 'juan@delco.cl',
                'disabled': False,
                'emailValidated': True,
                'groups': ['g1'],
            },
            {
                'id': 'u2',
                'emailAddress': 'otro@moreapp.com',
                'username': 'otro@moreapp.com',
                'disabled': True,
                'emailValidated': False,
                'groups': [],
            },
        ]
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('usuarios_moreapp'), {'query': 'juan'})
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('juan@delco.cl', body)
        self.assertNotIn('otro@moreapp.com', body)
        self.assertIn('Usuario MoreApp', body)
        self.assertIn('Usuario interno', body)
        self.assertIn('No se copian a la intranet', body)
        self.assertNotIn('vincular técnico', body.lower())
        self.assertNotIn('mapear usuario', body.lower())
        self.assertIn('navUsuariosMoreapp-', body)
        self.assertIn('Técnicos', body)
        self.assertIn('Certelec Técnicos', body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_catalogo_views.list_datasource_entry_rows')
    @mock.patch('integraciones.moreapp_catalogo_views.list_datasources')
    def test_usuarios_moreapp_muestra_tecnicos_y_certelec(self, mocked_ds, mocked_rows):
        mocked_ds.return_value = [
            {
                'id': '690cdeaea4d2441300ec961d',
                'name': 'Tecnicos',
                'columnMapping': [{'id': 'id', 'label': 'Id'}, {'id': 'NOMBRES', 'label': 'NOMBRES'}],
            },
            {
                'id': '68cab36aca9ce33084120ae2',
                'name': 'Certelec Técnicos',
                'columnMapping': [{'id': 'id', 'label': 'Id'}, {'id': 'nombres', 'label': 'Nombres'}],
            },
        ]
        mocked_rows.return_value = {
            'total': 2,
            'columns': [{'id': 'id', 'label': 'Id'}, {'id': 'NOMBRES', 'label': 'NOMBRES'}],
            'rows': [
                {'id': '1', 'NOMBRES': 'ALEXANDER ANDRES GONTHIER'},
                {'id': '2', 'NOMBRES': 'ASHLY LLANQUIMAN'},
            ],
            'fetched': 2,
            'truncated': False,
        }
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('usuarios_moreapp_tecnicos'))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('ALEXANDER ANDRES GONTHIER', body)
        self.assertIn('ASHLY LLANQUIMAN', body)
        self.assertNotIn('vincular técnico', body.lower())
        mocked_rows.return_value = {
            'total': 1,
            'columns': [{'id': 'id', 'label': 'Id'}, {'id': 'nombres', 'label': 'Nombres'}],
            'rows': [{'id': '1', 'nombres': 'Juan Andres Maldonado Perez'}],
            'fetched': 1,
            'truncated': False,
        }
        certelec = self.client.get(reverse('usuarios_moreapp_certelec'))
        self.assertEqual(certelec.status_code, 200)
        certelec_body = certelec.content.decode('utf-8')
        self.assertIn('Juan Andres Maldonado Perez', certelec_body)
        self.assertIn('Certelec Técnicos', certelec_body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_registros_views.get_form_version')
    @mock.patch('integraciones.moreapp_registros_views.get_submission')
    @mock.patch('integraciones.moreapp_registros_views.get_form')
    def test_detalle_conserva_filtros_al_volver(self, mocked_form, mocked_sub, mocked_version):
        mocked_form.return_value = {
            'id': 'abc123abc123abc123abc123',
            'name': 'Reposición de Suministro',
            'status': 'ACTIVE',
            'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
        }
        mocked_version.return_value = {'fields': []}
        mocked_sub.return_value = {
            'id': 'sub123sub123sub123sub123',
            'info': {'userName': 'Juan Perez', 'date': '2026-10-14T12:00:00Z'},
            'data': {'numCliente': '858678'},
        }
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(
            reverse(
                'moreapp_api_registro',
                args=['abc123abc123abc123abc123', 'sub123sub123sub123sub123'],
            ),
            {'tecnico': 'Juan', 'num_cliente': '858678', 'fecha_desde': '2026-10-01'},
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('tecnico=Juan', body)
        self.assertIn('num_cliente=858678', body)
        self.assertIn('fecha_desde=2026-10-01', body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_registros_views.get_form_version')
    @mock.patch('integraciones.moreapp_registros_views.list_submissions')
    @mock.patch('integraciones.moreapp_registros_views.get_form')
    def test_registros_exporta_excel(self, mocked_form, mocked_subs, mocked_version):
        mocked_form.return_value = {
            'id': 'abc123abc123abc123abc123',
            'name': 'Reposición de Suministro',
            'status': 'ACTIVE',
            'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
        }
        mocked_version.return_value = {'fields': []}
        mocked_subs.return_value = {
            'totalSize': 1,
            'elements': [
                {
                    'id': 'sub123sub123sub123sub123',
                    'info': {'userName': 'Juan Perez', 'date': '2026-10-14T12:00:00Z'},
                    'data': {'numCliente': 858678, 'tecnicoResponsable': {'NOMBRES': 'Juan Perez'}},
                }
            ],
        }
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(
            reverse('moreapp_api_form_registros', args=['abc123abc123abc123abc123']),
            {'tecnico': 'Juan', 'export': 'xlsx'},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertIn(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            resp['Content-Type'],
        )
        self.assertIn('attachment', resp['Content-Disposition'])
        self.assertGreater(len(resp.content), 100)

    def test_usuarios_internos_distinguen_moreapp(self):
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('usuarios_list'))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('Usuario MoreApp', body)
        self.assertIn('Usuario interno', body)
        self.assertIn('Interno', body)
        self.assertIn('/usuarios/moreapp/', body)
        self.assertIn('data-bs-toggle="collapse"', body)
        self.assertIn('navUsuarios-', body)
        self.assertIn('navUsuariosMoreapp-', body)
        self.assertIn('Técnicos', body)
        self.assertIn('Certelec Técnicos', body)
        self.assertIn('/usuarios/moreapp/tecnicos/', body)
        self.assertIn('/usuarios/moreapp/certelec/', body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_catalogo_views.get_datasource_entries')
    @mock.patch('integraciones.moreapp_catalogo_views.get_datasource')
    @mock.patch('integraciones.moreapp_catalogo_views.list_datasources')
    def test_guardar_fuente_es_opt_in(self, mocked_list, mocked_get, mocked_entries):
        from integraciones.models import MoreAppFuenteInterna

        mocked_list.return_value = [
            {
                'id': '6491e945a2abd076df165abc',
                'name': 'comunas',
                'searchType': 'LOCAL_FALLBACK',
                'enabled': True,
                'columnMapping': [{'id': 'id', 'label': 'Id'}, {'id': 'comuna', 'label': 'Comuna'}],
            },
            {
                'id': '690cdeaea4d2441300ec961d',
                'name': 'Tecnicos',
                'searchType': 'LOCAL_FALLBACK',
                'enabled': True,
                'columnMapping': [{'id': 'id', 'label': 'Id'}, {'id': 'NOMBRES', 'label': 'NOMBRES'}],
            },
        ]
        mocked_get.return_value = mocked_list.return_value[0]
        mocked_entries.return_value = {'numberOfEntries': 33, 'dataKeyList': mocked_list.return_value[0]['columnMapping']}
        self.client.login(username=self.admin.rut, password=self.password)
        save = self.client.post(
            reverse('moreapp_api_datasource_guardar', args=['6491e945a2abd076df165abc'])
        )
        self.assertEqual(save.status_code, 302)
        self.assertTrue(
            MoreAppFuenteInterna.objects.filter(datasource_id='6491e945a2abd076df165abc').exists()
        )
        self.assertEqual(MoreAppFuenteInterna.objects.count(), 1)
        listing = self.client.get(reverse('moreapp_api_datasources'))
        body = listing.content.decode('utf-8')
        self.assertIn('Guardada', body)
        self.assertIn('No se guarda', body)
        self.assertIn('Guardar en intranet', body)
        only_local = self.client.get(reverse('moreapp_api_datasources'), {'solo_internas': '1'})
        only_body = only_local.content.decode('utf-8')
        self.assertIn('comunas', only_body)
        self.assertIn('Tecnicos', only_body)
        self.assertIn('data-intranet="1"', only_body)
        self.assertIn('js-ds-card', only_body)
        self.assertIn('data-ubicacion="intranet"', only_body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_catalogo_views.get_datasource_entries')
    @mock.patch('integraciones.moreapp_catalogo_views.get_datasource')
    def test_importar_filas_solo_si_esta_guardada(self, mocked_get, mocked_entries):
        from integraciones.models import MoreAppFuenteFila, MoreAppFuenteInterna

        mocked_get.return_value = {
            'id': '6491e945a2abd076df165abc',
            'name': 'comunas',
            'columnMapping': [{'id': 'id', 'label': 'Id'}, {'id': 'comuna', 'label': 'Comuna'}],
        }
        mocked_entries.return_value = {'numberOfEntries': 2, 'dataKeyList': mocked_get.return_value['columnMapping']}
        self.client.login(username=self.admin.rut, password=self.password)
        archivo = SimpleUploadedFile(
            'comunas.json',
            json.dumps([{'id': 'stgo', 'comuna': 'Santiago'}, {'id': 'maipu', 'comuna': 'Maipu'}]).encode('utf-8'),
            content_type='application/json',
        )
        denied = self.client.post(
            reverse('moreapp_api_datasource_importar', args=['6491e945a2abd076df165abc']),
            {'archivo': archivo},
        )
        self.assertEqual(denied.status_code, 302)
        self.assertEqual(MoreAppFuenteFila.objects.count(), 0)
        self.client.post(reverse('moreapp_api_datasource_guardar', args=['6491e945a2abd076df165abc']))
        archivo2 = SimpleUploadedFile(
            'comunas.json',
            json.dumps([{'id': 'stgo', 'comuna': 'Santiago'}, {'id': 'maipu', 'comuna': 'Maipu'}]).encode('utf-8'),
            content_type='application/json',
        )
        imported = self.client.post(
            reverse('moreapp_api_datasource_importar', args=['6491e945a2abd076df165abc']),
            {'archivo': archivo2},
        )
        self.assertEqual(imported.status_code, 302)
        fuente = MoreAppFuenteInterna.objects.get(datasource_id='6491e945a2abd076df165abc')
        self.assertEqual(fuente.filas.count(), 2)
        self.assertEqual(MoreAppFuenteFila.objects.get(entry_id='stgo').valores.get('comuna'), 'Santiago')
        detail = self.client.get(reverse('moreapp_api_datasource', args=['6491e945a2abd076df165abc']))
        self.assertIn('Santiago', detail.content.decode('utf-8'))
        self.assertIn('Copia en la intranet', detail.content.decode('utf-8'))

    def test_listado_formularios_requiere_login(self):
        resp = self.client.get(reverse('moreapp_formularios_list'))
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/login', resp.url)

    def test_tecnico_ve_formularios_en_sidebar(self):
        self.client.login(username=self.tecnico.rut, password=self.password)
        resp = self.client.get(reverse('dashboard'))
        body = resp.content.decode('utf-8')
        self.assertIn('Formularios asignados', body)
        self.assertNotIn('Pendientes MoreApp', body)
        self.assertNotIn('Catálogo MoreApp', body)
        self.assertNotIn('Tareas MoreApp', body)

    def test_tecnico_sin_asignacion_lista_vacia(self):
        self.client.login(username=self.tecnico.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_formularios_list'))
        self.assertEqual(resp.status_code, 200)
        self.assertIn('No tienes formularios asignados', resp.content.decode('utf-8'))

    def test_tecnico_no_asigna_formularios(self):
        self.client.login(username=self.tecnico.rut, password=self.password)
        resp = self.client.get(
            reverse('moreapp_formulario_asignar', args=['abc123abc123abc123abc123'])
        )
        self.assertEqual(resp.status_code, 403)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_formularios_views.list_all_forms')
    @mock.patch('integraciones.moreapp_catalogo_views.list_all_forms')
    def test_admin_lista_y_asigna_formulario(self, mocked_catalogo_forms, mocked_list_forms):
        from integraciones.models import MoreAppFormularioAsignado

        form = {
            'id': 'abc123abc123abc123abc123',
            'name': 'Instalacion',
            'status': 'ACTIVE',
            'publishedVersion': {'formVersion': '1'},
        }
        extra = {
            'id': 'def456def456def456def456',
            'name': 'Retiro',
            'status': 'ACTIVE',
            'publishedVersion': {'formVersion': '1'},
        }
        mocked_catalogo_forms.return_value = [form, extra]
        mocked_list_forms.return_value = [form, extra]
        adminis = Usuario.objects.create_user(
            rut='66666666-6',
            email='adm-asig@delco.cl',
            password=self.password,
            nombre='Adm',
            apellido='Asig',
            nombre_interno='adm_api',
            rol='ADMINISTRATIVO',
            is_active=True,
        )
        self.client.login(username=self.admin.rut, password=self.password)
        vacio = self.client.get(reverse('moreapp_formularios_list'))
        self.assertEqual(vacio.status_code, 200)
        self.assertIn('Aún no hay formularios asignados', vacio.content.decode('utf-8'))
        catalogo = self.client.get(reverse('moreapp_api_catalogo'))
        self.assertEqual(catalogo.status_code, 200)
        body = catalogo.content.decode('utf-8')
        self.assertIn('Instalacion', body)
        self.assertIn('Asignar usuarios', body)
        self.assertIn('modalAsignarFormulario', body)
        self.assertIn('type="checkbox"', body)
        self.assertIn('adm_api', body)
        self.assertNotIn('tec_api', body)
        get_page = self.client.get(
            reverse('moreapp_formulario_asignar', args=['abc123abc123abc123abc123'])
        )
        self.assertEqual(get_page.status_code, 302)
        self.assertEqual(get_page.url, reverse('moreapp_api_catalogo'))
        denied = self.client.post(
            reverse('moreapp_formulario_asignar', args=['abc123abc123abc123abc123']),
            {'nombre': 'Instalacion', 'usuarios': [str(self.tecnico.pk)]},
        )
        self.assertEqual(denied.status_code, 302)
        self.assertFalse(
            MoreAppFormularioAsignado.objects.filter(
                usuario=self.tecnico,
                form_id='abc123abc123abc123abc123',
            ).exists()
        )
        MoreAppFormularioAsignado.objects.create(
            usuario=self.tecnico,
            form_id='abc123abc123abc123abc123',
            form_name='Instalacion',
            asignado_por=self.admin,
        )
        post = self.client.post(
            reverse('moreapp_formulario_asignar', args=['abc123abc123abc123abc123']),
            {'nombre': 'Instalacion', 'usuarios': [str(adminis.pk), str(self.admin.pk)]},
        )
        self.assertEqual(post.status_code, 302)
        self.assertEqual(
            MoreAppFormularioAsignado.objects.filter(
                form_id='abc123abc123abc123abc123',
            ).count(),
            3,
        )
        self.assertTrue(
            MoreAppFormularioAsignado.objects.filter(
                usuario=self.tecnico,
                form_id='abc123abc123abc123abc123',
            ).exists()
        )
        self.assertTrue(
            MoreAppFormularioAsignado.objects.filter(
                usuario=adminis,
                form_id='abc123abc123abc123abc123',
            ).exists()
        )
        listing = self.client.get(reverse('moreapp_formularios_list'))
        self.assertEqual(listing.status_code, 200)
        list_body = listing.content.decode('utf-8')
        self.assertIn('Instalacion', list_body)
        self.assertNotIn('Retiro', list_body)
        self.assertIn('adm_api', list_body)
        self.assertNotIn('Asignar usuarios', list_body)
        self.assertNotIn('Ir al catálogo', list_body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_formularios_views.list_all_forms')
    def test_tecnico_solo_ve_formularios_asignados(self, mocked_forms):
        from integraciones.models import MoreAppFormularioAsignado

        mocked_forms.return_value = [
            {
                'id': 'abc123abc123abc123abc123',
                'name': 'Instalacion',
                'status': 'ACTIVE',
                'publishedVersion': {'formVersion': '1'},
            },
            {
                'id': 'def456def456def456def456',
                'name': 'Retiro',
                'status': 'ACTIVE',
                'publishedVersion': {'formVersion': '1'},
            },
        ]
        MoreAppFormularioAsignado.objects.create(
            usuario=self.tecnico,
            form_id='abc123abc123abc123abc123',
            form_name='Instalacion',
        )
        self.client.login(username=self.tecnico.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_formularios_list'))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('Instalacion', body)
        self.assertNotIn('Retiro', body)
        self.assertNotIn('Asignar usuarios', body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_registros_views.get_form_version')
    @mock.patch('integraciones.moreapp_registros_views.list_submissions')
    @mock.patch('integraciones.moreapp_registros_views.get_form')
    def test_tecnico_con_asignacion_accede_registros(self, mocked_form, mocked_subs, mocked_version):
        from integraciones.models import MoreAppFormularioAsignado

        MoreAppFormularioAsignado.objects.create(
            usuario=self.tecnico,
            form_id='abc123abc123abc123abc123',
            form_name='Instalacion',
        )
        mocked_form.return_value = {
            'id': 'abc123abc123abc123abc123',
            'name': 'Instalacion',
            'status': 'ACTIVE',
            'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
        }
        mocked_version.return_value = {'id': 'ver123ver123ver123ver123', 'fields': []}
        mocked_subs.return_value = {'data': []}
        self.client.login(username=self.tecnico.rut, password=self.password)
        resp = self.client.get(
            reverse('moreapp_api_form_registros', args=['abc123abc123abc123abc123'])
        )
        self.assertEqual(resp.status_code, 200)
        self.assertIn('Instalacion', resp.content.decode('utf-8'))

    def test_tareas_requiere_login(self):
        resp = self.client.get(reverse('moreapp_tareas_list'))
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/login', resp.url)

    def test_listado_tareas_ofrece_carga_masiva(self):
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_tareas_list'))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('modalLoteCsv', body)
        self.assertIn('modalLoteClientes', body)
        self.assertIn('modalLoteListado', body)
        self.assertIn('data-bs-target="#modalLoteCsv"', body)
        self.assertIn('data-bs-target="#modalLoteClientes"', body)
        self.assertIn('data-bs-target="#modalLoteListado"', body)
        self.assertIn(reverse('moreapp_tareas_lote_csv'), body)
        self.assertIn(reverse('moreapp_tareas_lote_clientes'), body)
        self.assertIn(reverse('moreapp_tareas_lote_listado'), body)

    def test_tecnico_no_accede_tareas(self):
        self.client.login(username=self.tecnico.rut, password=self.password)
        self.assertEqual(self.client.get(reverse('moreapp_tareas_list')).status_code, 403)
        self.assertEqual(self.client.get(reverse('moreapp_tarea_crear')).status_code, 403)
        self.assertEqual(self.client.get(reverse('moreapp_tarea_detalle', args=[1])).status_code, 403)
        self.assertEqual(self.client.get(reverse('moreapp_tarea_pdf', args=[1])).status_code, 403)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_tareas_views.list_users')
    @mock.patch('integraciones.moreapp_tareas_views.list_all_forms')
    def test_admin_ve_formulario_de_tarea(self, mocked_forms, mocked_users):
        mocked_forms.return_value = [
            {
                'id': 'abc123abc123abc123abc123',
                'name': 'Instalacion',
                'status': 'ACTIVE',
                'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
            }
        ]
        mocked_users.return_value = [
            {
                'id': 'u1',
                'emailAddress': 'campo@more.app',
                'disabled': False,
                'settings': {'displayName': 'Tecnico More'},
                'groups': [{'id': 'g1', 'name': 'Terreno'}],
            }
        ]
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_tarea_crear'))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('Instalacion', body)
        self.assertIn('campo@more.app', body)
        self.assertIn('Terreno', body)
        self.assertIn('Al momento', body)
        self.assertIn('Fecha y hora programadas', body)
        self.assertIn('Pin, firma, catálogo, dibujo y archivo no se pueden precargar', body)
        self.assertNotIn('No se pudieron cargar grupos', body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_tareas_views.list_users')
    @mock.patch('integraciones.moreapp_tareas_views.list_all_forms')
    def test_tarea_omite_formulario_sin_version_publicada(self, mocked_forms, mocked_users):
        mocked_forms.return_value = [
            {
                'id': '6a074c42668c2d11d36cfcf5',
                'name': 'Borrador sin publicar',
                'status': 'ACTIVE',
                'publishedVersion': {'formVersion': ''},
            },
            {
                'id': 'abc123abc123abc123abc123',
                'name': 'Instalacion',
                'status': 'ACTIVE',
                'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
            },
        ]
        mocked_users.return_value = [
            {'id': 'u1', 'emailAddress': 'campo@more.app', 'disabled': False}
        ]
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_tarea_crear'))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('Instalacion', body)
        self.assertNotIn('Borrador sin publicar', body)
        blocked = self.client.post(
            reverse('moreapp_tarea_crear'),
            {
                'form_id': '6a074c42668c2d11d36cfcf5',
                'mensaje': 'Ir a terreno',
                'envio': 'IMMEDIATE',
                'usuarios': ['campo@more.app'],
            },
        )
        self.assertEqual(blocked.status_code, 200)
        self.assertIn('versión publicada', blocked.content.decode('utf-8'))

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_tareas_views.group_names_by_id')
    @mock.patch('integraciones.moreapp_tareas_views.list_users')
    @mock.patch('integraciones.moreapp_tareas_views.list_all_forms')
    def test_grupos_muestran_nombre_oficial_moreapp(self, mocked_forms, mocked_users, mocked_names):
        mocked_forms.return_value = [
            {
                'id': 'abc123abc123abc123abc123',
                'name': 'Instalacion',
                'status': 'ACTIVE',
                'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
            }
        ]
        mocked_users.return_value = [
            {
                'id': 'u1',
                'emailAddress': 'campo@more.app',
                'disabled': False,
                'groups': ['65f0a25a128983033c9fd314'],
            }
        ]
        mocked_names.return_value = {'65f0a25a128983033c9fd314': 'Cuadrilla Norte'}
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_tarea_crear'))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('Cuadrilla Norte', body)
        self.assertNotIn('Grupo 9fd314', body)
        self.assertNotIn('Guardar nombres', body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_api.list_users')
    def test_grupos_moreapp_salen_de_usuarios(self, mocked_users):
        from integraciones.moreapp_api import list_groups

        mocked_users.return_value = [
            {'id': 'u1', 'groups': [{'id': 'g1', 'name': 'Terreno'}]},
            {'id': 'u2', 'groups': [{'id': 'g1', 'name': 'Terreno'}, {'id': 'g2', 'name': 'Oficina'}]},
        ]
        groups = list_groups()
        self.assertEqual({item['id'] for item in groups}, {'g1', 'g2'})
        self.assertEqual({item['name'] for item in groups}, {'Terreno', 'Oficina'})

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_tareas_views.create_task')
    @mock.patch('integraciones.moreapp_tareas_views._campos_desde_formulario')
    @mock.patch('integraciones.moreapp_tareas_views.get_form')
    @mock.patch('integraciones.moreapp_tareas_views.list_users')
    @mock.patch('integraciones.moreapp_tareas_views.list_all_forms')
    def test_admin_crea_tarea_inmediata(
        self,
        mocked_forms,
        mocked_users,
        mocked_get_form,
        mocked_campos,
        mocked_create,
    ):
        from integraciones.models import MoreAppTareaEnviada

        form = {
            'id': 'abc123abc123abc123abc123',
            'name': 'Instalacion',
            'status': 'ACTIVE',
            'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
        }
        mocked_forms.return_value = [form]
        mocked_users.return_value = [
            {'id': 'u1', 'emailAddress': 'campo@more.app', 'disabled': False}
        ]
        mocked_get_form.return_value = form
        mocked_campos.return_value = [
            {'widget': 'text', 'data_name': 'cliente', 'label': 'Cliente'},
            {'widget': 'signature', 'data_name': 'firma', 'label': 'Firma'},
            {'widget': 'file', 'data_name': 'archivo', 'label': 'Archivo'},
        ]
        mocked_create.return_value = {'id': 'task-99'}
        self.client.login(username=self.admin.rut, password=self.password)
        sin_dest = self.client.post(
            reverse('moreapp_tarea_crear'),
            {
                'form_id': 'abc123abc123abc123abc123',
                'mensaje': 'Ir a terreno',
                'envio': 'IMMEDIATE',
            },
        )
        self.assertEqual(sin_dest.status_code, 200)
        self.assertFalse(MoreAppTareaEnviada.objects.exists())
        programada_sin_fecha = self.client.post(
            reverse('moreapp_tarea_crear'),
            {
                'form_id': 'abc123abc123abc123abc123',
                'mensaje': 'Ir a terreno',
                'envio': 'SCHEDULED',
                'usuarios': ['campo@more.app'],
            },
        )
        self.assertEqual(programada_sin_fecha.status_code, 200)
        self.assertFalse(MoreAppTareaEnviada.objects.exists())
        resp = self.client.post(
            reverse('moreapp_tarea_crear'),
            {
                'form_id': 'abc123abc123abc123abc123',
                'mensaje': 'Instalar medidor',
                'envio': 'IMMEDIATE',
                'usuarios': ['campo@more.app'],
                'ubicacion': 'Av. Principal 100',
                'latitud': '-33.45',
                'longitud': '-70.66',
                'prefill__cliente': 'ACME',
                'prefill__firma': 'no-debe-irse',
                'prefill__archivo': 'tampoco',
            },
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp.url, reverse('moreapp_tareas_list'))
        mocked_create.assert_called_once()
        form_id, payload = mocked_create.call_args[0]
        self.assertEqual(form_id, 'abc123abc123abc123abc123')
        self.assertEqual(payload['recipients'], ['campo@more.app'])
        self.assertEqual(payload['message'], 'Instalar medidor')
        self.assertEqual(payload['publishInfo']['type'], 'IMMEDIATE')
        self.assertEqual(payload['data'], {'cliente': 'ACME'})
        self.assertEqual(payload['location']['formattedValue'], 'Av. Principal 100')
        self.assertEqual(payload['location']['coordinates']['latitude'], -33.45)
        self.assertEqual(payload['location']['coordinates']['longitude'], -70.66)
        tarea = MoreAppTareaEnviada.objects.get()
        self.assertEqual(tarea.task_id, 'task-99')
        self.assertEqual(tarea.mensaje, 'Instalar medidor')
        self.assertEqual(tarea.datos_precarga, {'cliente': 'ACME'})
        self.assertEqual(tarea.estado, 'ASIGNADA')
        self.assertEqual(tarea.cuenta, 'ACME')
        self.assertTrue(tarea.eventos.filter(tipo='CREADA').exists())
        self.assertTrue(tarea.eventos.filter(tipo='ASIGNADA').exists())
        self.assertTrue(tarea.alertas.filter(tipo='CREACION', usuario=self.admin, broadcast=False).exists())
        self.assertTrue(tarea.alertas.filter(tipo='CREACION', broadcast=True).exists())
        listing = self.client.get(reverse('moreapp_tareas_list'))
        self.assertEqual(listing.status_code, 200)
        list_body = listing.content.decode('utf-8')
        self.assertIn('Instalar medidor', list_body)
        self.assertIn('aviso llega en MoreApp, no por correo', list_body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_tareas_views.create_task')
    @mock.patch('integraciones.moreapp_tareas_views._campos_desde_formulario')
    @mock.patch('integraciones.moreapp_tareas_views.get_form')
    @mock.patch('integraciones.moreapp_tareas_views.list_users')
    @mock.patch('integraciones.moreapp_tareas_views.list_all_forms')
    def test_tarea_a_grupo_expande_correos_y_no_manda_ids(
        self,
        mocked_forms,
        mocked_users,
        mocked_get_form,
        mocked_campos,
        mocked_create,
    ):
        form = {
            'id': 'abc123abc123abc123abc123',
            'name': 'Instalacion',
            'status': 'ACTIVE',
            'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
        }
        mocked_forms.return_value = [form]
        mocked_users.return_value = [
            {
                'id': 'u1',
                'emailAddress': 'campo@more.app',
                'disabled': False,
                'groups': ['g1'],
            },
            {
                'id': 'u2',
                'emailAddress': 'oficina@more.app',
                'disabled': False,
                'groups': ['g1', 'g2'],
            },
            {
                'id': 'u3',
                'emailAddress': 'otro@more.app',
                'disabled': False,
                'groups': ['g2'],
            },
        ]
        mocked_get_form.return_value = form
        mocked_campos.return_value = []
        mocked_create.return_value = {'id': 'task-group-1'}
        self.client.login(username=self.admin.rut, password=self.password)
        crear = self.client.get(reverse('moreapp_tarea_crear'))
        self.assertEqual(crear.status_code, 200)
        crear_body = crear.content.decode('utf-8')
        self.assertIn('2 personas', crear_body)
        self.assertIn('campo@more.app', crear_body)
        self.assertNotIn('Grupo g1', crear_body)
        self.assertNotIn('Grupo 9fd314', crear_body)
        resp = self.client.post(
            reverse('moreapp_tarea_crear'),
            {
                'form_id': 'abc123abc123abc123abc123',
                'mensaje': 'Ir a terreno',
                'envio': 'IMMEDIATE',
                'grupos': ['g1'],
            },
        )
        self.assertEqual(resp.status_code, 302)
        mocked_create.assert_called_once()
        _form_id, payload = mocked_create.call_args[0]
        self.assertEqual(sorted(payload['recipients']), ['campo@more.app', 'oficina@more.app'])
        self.assertNotIn('g1', payload['recipients'])
        self.assertNotIn('groups', payload)
        self.assertEqual(payload['data'], {})
        self.assertNotIn('location', payload)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_tareas_views._campos_desde_formulario')
    @mock.patch('integraciones.moreapp_tareas_views.get_form')
    def test_campos_precarga_omite_firma_y_archivo(self, mocked_get_form, mocked_campos):
        mocked_get_form.return_value = {'id': 'abc123abc123abc123abc123'}
        mocked_campos.return_value = [
            {'widget': 'text', 'data_name': 'cliente', 'label': 'Cliente'},
            {'widget': 'pin', 'data_name': 'pin', 'label': 'Pin'},
            {'widget': 'signature', 'data_name': 'firma', 'label': 'Firma'},
            {'widget': 'catalog', 'data_name': 'catalogo', 'label': 'Catalogo'},
            {'widget': 'drawing', 'data_name': 'dibujo', 'label': 'Dibujo'},
            {'widget': 'file', 'data_name': 'archivo', 'label': 'Archivo'},
        ]
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(
            reverse('moreapp_tarea_campos', args=['abc123abc123abc123abc123'])
        )
        self.assertEqual(resp.status_code, 200)
        payload = resp.json()
        self.assertTrue(payload['ok'])
        keys = [item['key'] for item in payload['campos']]
        self.assertEqual(keys, ['cliente'])

    def test_tecnico_no_accede_lote_tareas(self):
        self.client.login(username=self.tecnico.rut, password=self.password)
        self.assertEqual(self.client.get(reverse('moreapp_tareas_lote_csv')).status_code, 403)
        self.assertEqual(self.client.get(reverse('moreapp_tareas_lote_clientes')).status_code, 403)
        self.assertEqual(self.client.get(reverse('moreapp_tareas_lote_listado')).status_code, 403)
        self.assertEqual(self.client.get(reverse('moreapp_tareas_lote_plantilla')).status_code, 403)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    def test_plantilla_csv_tareas(self):
        import io
        from openpyxl import load_workbook

        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_tareas_lote_plantilla'))
        self.assertEqual(resp.status_code, 200)
        self.assertIn('attachment; filename="tareas_moreapp.xlsx"', resp['Content-Disposition'])
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        self.assertEqual(
            [ws.cell(1, col).value for col in range(1, 7)],
            ['formulario', 'destinatario', 'mensaje', 'fecha', 'ubicacion', 'cliente'],
        )
        self.assertIsNone(ws.cell(2, 1).value)
        self.assertNotEqual(ws.cell(1, 1).value, ws.cell(1, 2).value)

    def test_lote_csv_abre_popup_del_listado(self):
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_tareas_lote_csv'))
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse('moreapp_tareas_list'), resp.url)
        self.assertIn('lote=csv', resp.url)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_tareas_views.list_datasource_entry_rows')
    @mock.patch('integraciones.moreapp_tareas_views._campos_desde_formulario')
    @mock.patch('integraciones.moreapp_tareas_views.get_form')
    @mock.patch('integraciones.moreapp_tareas_views.create_task')
    @mock.patch('integraciones.moreapp_tareas_views.list_users')
    @mock.patch('integraciones.moreapp_tareas_views.list_all_forms')
    def test_lote_csv_valida_y_crea_tareas(
        self,
        mocked_forms,
        mocked_users,
        mocked_create,
        mocked_get_form,
        mocked_campos,
        mocked_rows,
    ):
        from integraciones.models import MoreAppTareaEnviada

        form = {
            'id': 'abc123abc123abc123abc123',
            'name': 'Instalacion',
            'status': 'ACTIVE',
            'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
        }
        mocked_forms.return_value = [form]
        mocked_users.return_value = [
            {
                'id': 'u1',
                'emailAddress': 'campo@more.app',
                'disabled': False,
                'settings': {'displayName': 'Tecnico More'},
                'groups': [{'id': 'g1', 'name': 'Terreno'}],
            }
        ]
        mocked_get_form.return_value = form
        mocked_campos.return_value = [
            {
                'widget': 'search',
                'data_name': 'cliente1',
                'label': 'Cliente',
                'datasource_id': 'dscliente1dscliente1ds01',
            }
        ]
        mocked_rows.return_value = {
            'rows': [{'CUENTA': '1001', 'DIRECCION': 'Calle Uno', 'id': 'row-1'}]
        }
        mocked_create.side_effect = [{'id': 'task-csv-1'}, {'id': 'task-csv-2'}]
        csv_body = (
            'formulario,destinatario,mensaje,fecha,ubicacion,cliente\n'
            'Instalacion,campo@more.app,Visita,2026-10-23 09:00,,1001\n'
            'Instalacion,Terreno,Visita grupo,,,\n'
            'No existe,campo@more.app,Mal,,,,\n'
        )
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.post(
            reverse('moreapp_tareas_lote_csv'),
            {
                'archivo': SimpleUploadedFile(
                    'lote.csv',
                    csv_body.encode('utf-8'),
                    content_type='text/csv',
                )
            },
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(mocked_create.call_count, 2)
        primera = mocked_create.call_args_list[0][0][1]
        self.assertEqual(primera['recipients'], ['campo@more.app'])
        self.assertEqual(primera['data']['cliente1']['CUENTA'], '1001')
        segunda = mocked_create.call_args_list[1][0][1]
        self.assertEqual(segunda['recipients'], ['campo@more.app'])
        self.assertEqual(segunda['data'], {})
        self.assertEqual(MoreAppTareaEnviada.objects.filter(error='').count(), 2)
        body = resp.content.decode('utf-8')
        self.assertIn('Formulario no encontrado', body)
        self.assertIn('modalLoteCsv', body)
        self.assertIn("var open = 'csv'", body)

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_tareas_views.list_datasource_entry_rows')
    @mock.patch('integraciones.moreapp_tareas_views._campos_desde_formulario')
    @mock.patch('integraciones.moreapp_tareas_views.get_form')
    @mock.patch('integraciones.moreapp_tareas_views.create_task')
    @mock.patch('integraciones.moreapp_tareas_views.list_users')
    @mock.patch('integraciones.moreapp_tareas_views.list_all_forms')
    def test_lote_desde_clientes_una_tarea_por_cliente(
        self,
        mocked_forms,
        mocked_users,
        mocked_create,
        mocked_get_form,
        mocked_campos,
        mocked_rows,
    ):
        from clientes.models import Cliente
        from integraciones.models import MoreAppTareaEnviada

        form = {
            'id': 'abc123abc123abc123abc123',
            'name': 'Instalacion',
            'status': 'ACTIVE',
            'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
        }
        mocked_forms.return_value = [form]
        mocked_users.return_value = [
            {
                'id': 'u1',
                'emailAddress': 'campo@more.app',
                'disabled': False,
                'settings': {'displayName': 'Tecnico More'},
                'groups': [{'id': 'g1', 'name': 'Terreno'}],
            },
            {
                'id': 'u2',
                'emailAddress': 'oficina@more.app',
                'disabled': False,
                'groups': ['g1'],
            },
        ]
        mocked_get_form.return_value = form
        mocked_campos.return_value = [
            {
                'widget': 'search',
                'data_name': 'cliente1',
                'label': 'Cliente',
                'datasource_id': 'dscliente1dscliente1ds01',
            }
        ]
        mocked_rows.return_value = {
            'rows': [{'CUENTA': '1001', 'DIRECCION': 'Calle Uno', 'COMUNA': 'Santiago', 'id': 'row-1'}]
        }
        mocked_create.side_effect = [{'id': 'task-cli-1'}, {'id': 'task-cli-2'}]
        coincidente = Cliente.objects.create(
            numero_cliente='1001',
            direccion='Calle Uno',
            comuna='Santiago',
        )
        sin_match = Cliente.objects.create(
            numero_cliente='2002',
            direccion='Calle Dos',
            comuna='Providencia',
        )
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.post(
            reverse('moreapp_tareas_lote_clientes'),
            {
                'form_id': 'abc123abc123abc123abc123',
                'mensaje': 'Visita masiva',
                'grupos': ['g1'],
                'clientes': [str(coincidente.pk), str(sin_match.pk)],
            },
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(mocked_create.call_count, 2)
        payloads = [call[0][1] for call in mocked_create.call_args_list]
        self.assertEqual(payloads[0]['recipients'], ['campo@more.app', 'oficina@more.app'])
        self.assertEqual(payloads[1]['recipients'], ['campo@more.app', 'oficina@more.app'])
        cuentas = [
            (item['data'] or {}).get('cliente1', {}).get('CUENTA')
            for item in payloads
        ]
        self.assertEqual(sorted([item for item in cuentas if item]), ['1001'])
        self.assertEqual(MoreAppTareaEnviada.objects.filter(error='').count(), 2)
        mensajes = list(
            MoreAppTareaEnviada.objects.values_list('mensaje', flat=True)
        )
        self.assertTrue(any('1001' in item for item in mensajes))
        self.assertTrue(any('2002' in item for item in mensajes))

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_tareas_views.list_users')
    @mock.patch('integraciones.moreapp_tareas_views.list_all_forms')
    def test_lote_clientes_json_para_popup(self, mocked_forms, mocked_users):
        from clientes.models import Cliente

        mocked_forms.return_value = [
            {
                'id': 'abc123abc123abc123abc123',
                'name': 'Instalacion',
                'status': 'ACTIVE',
                'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
            }
        ]
        mocked_users.return_value = [
            {
                'id': 'u1',
                'emailAddress': 'campo@more.app',
                'disabled': False,
                'settings': {'displayName': 'Tecnico More'},
                'groups': [{'id': 'g1', 'name': 'Terreno'}],
            }
        ]
        Cliente.objects.create(numero_cliente='1001', direccion='Calle Uno', comuna='Santiago')
        self.client.login(username=self.admin.rut, password=self.password)
        redir = self.client.get(reverse('moreapp_tareas_lote_clientes'))
        self.assertEqual(redir.status_code, 302)
        self.assertIn('lote=clientes', redir.url)
        resp = self.client.get(reverse('moreapp_tareas_lote_clientes'), {'json': '1'})
        self.assertEqual(resp.status_code, 200)
        payload = resp.json()
        self.assertTrue(payload['ok'])
        self.assertEqual(payload['formularios'][0]['name'], 'Instalacion')
        self.assertEqual(payload['usuarios'][0]['email'], 'campo@more.app')
        self.assertEqual(payload['clientes'][0]['numero_cliente'], '1001')

    @override_settings(
        MOREAPP_API_ENABLED=True,
        MOREAPP_API_KEY='test-key',
        MOREAPP_CUSTOMER_ID='123',
    )
    @mock.patch('integraciones.moreapp_tareas_views.create_task')
    @mock.patch('integraciones.moreapp_tareas_views.list_users')
    @mock.patch('integraciones.moreapp_tareas_views.list_all_forms')
    def test_lote_listado_crea_cantidad_por_fila(self, mocked_forms, mocked_users, mocked_create):
        from integraciones.models import MoreAppTareaEnviada

        form = {
            'id': 'abc123abc123abc123abc123',
            'name': 'Instalacion',
            'status': 'ACTIVE',
            'publishedVersion': {'formVersion': 'ver123ver123ver123ver123'},
        }
        mocked_forms.return_value = [form]
        mocked_users.return_value = [
            {
                'id': 'u1',
                'emailAddress': 'campo@more.app',
                'disabled': False,
                'settings': {'displayName': 'Tecnico More'},
                'groups': [{'id': 'g1', 'name': 'Terreno'}],
            },
            {
                'id': 'u2',
                'emailAddress': 'oficina@more.app',
                'disabled': False,
                'groups': ['g1'],
            },
        ]
        mocked_create.side_effect = [{'id': 't1'}, {'id': 't2'}, {'id': 't3'}]
        self.client.login(username=self.admin.rut, password=self.password)
        json_resp = self.client.get(reverse('moreapp_tareas_lote_listado'), {'json': '1'})
        self.assertEqual(json_resp.status_code, 200)
        self.assertTrue(json_resp.json()['ok'])
        resp = self.client.post(
            reverse('moreapp_tareas_lote_listado'),
            {
                'mensaje': 'Ronda de terreno',
                'form_id': ['abc123abc123abc123abc123', 'abc123abc123abc123abc123'],
                'destinatario': ['user:campo@more.app', 'group:g1'],
                'cantidad': ['2', '1'],
            },
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(mocked_create.call_count, 3)
        self.assertEqual(MoreAppTareaEnviada.objects.filter(error='').count(), 3)
        mensajes = list(MoreAppTareaEnviada.objects.values_list('mensaje', flat=True))
        self.assertTrue(any('1/2' in item for item in mensajes))
        self.assertTrue(any('2/2' in item for item in mensajes))
        payloads = [call[0][1] for call in mocked_create.call_args_list]
        self.assertEqual(payloads[0]['recipients'], ['campo@more.app'])
        self.assertEqual(sorted(payloads[2]['recipients']), ['campo@more.app', 'oficina@more.app'])

    def test_lote_listado_ignora_filas_vacias_extra(self):
        from integraciones.moreapp_tareas_lote import parse_filas_listado

        class Post(dict):
            def getlist(self, key):
                return list(self.get(key) or [])

        filas, err = parse_filas_listado(
            Post(
                {
                    'form_id': ['abc123abc123abc123abc123', '', ''],
                    'destinatario': ['user:campo@more.app', '', ''],
                    'cantidad': ['2', '1', '1'],
                }
            )
        )
        self.assertEqual(err, '')
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0]['cantidad'], 2)

    def test_lote_csv_cuenta_filas_sobre_el_maximo(self):
        from integraciones.moreapp_tareas_lote import MAX_LOTE, parse_lote_csv

        body = 'formulario,destinatario,mensaje,fecha,ubicacion,cliente\n'
        body += ''.join('F%s,a@b.c,m,,,,\n' % idx for idx in range(MAX_LOTE + 1))
        rows, total = parse_lote_csv(body.encode('utf-8'))
        self.assertEqual(total, MAX_LOTE + 1)
        self.assertEqual(len(rows), MAX_LOTE)
        self.assertEqual(rows[-1]['formulario'], 'F%s' % (MAX_LOTE - 1))

    def test_lote_xlsx_usa_hoja_tareas_si_instrucciones_activa(self):
        from io import BytesIO

        from openpyxl import load_workbook

        from integraciones.moreapp_tareas_lote import parse_lote_xlsx, plantilla_xlsx

        wb = load_workbook(BytesIO(plantilla_xlsx()))
        wb['Tareas']['A2'] = 'Instalacion'
        wb['Tareas']['B2'] = 'campo@more.app'
        wb.active = wb['Instrucciones']
        buf = BytesIO()
        wb.save(buf)
        rows, total = parse_lote_xlsx(buf.getvalue())
        self.assertEqual(total, 1)
        self.assertEqual(rows[0]['formulario'], 'Instalacion')
        self.assertEqual(rows[0]['destinatario'], 'campo@more.app')


class MoreAppTareaCicloTests(TestCase):
    def setUp(self):
        self.password = 'Seguro1234'
        self.admin = Usuario.objects.create_user(
            rut='55555555-5',
            email='admin-ciclo@delco.cl',
            password=self.password,
            nombre='Admin',
            apellido='Ciclo',
            nombre_interno='admin_ciclo',
            rol='ADMIN',
            is_active=True,
        )
        self.administrativo = Usuario.objects.create_user(
            rut='56565656-5',
            email='adm-ciclo@delco.cl',
            password=self.password,
            nombre='Adm',
            apellido='Ciclo',
            nombre_interno='adm_ciclo',
            rol='ADMINISTRATIVO',
            is_active=True,
        )
        self.tecnico = Usuario.objects.create_user(
            rut='57575757-5',
            email='tec-ciclo@delco.cl',
            password=self.password,
            nombre='Tec',
            apellido='Ciclo',
            nombre_interno='tec_ciclo',
            rol='TECNICO',
            is_active=True,
        )

    def _tarea(self, **kwargs):
        from integraciones.models import MoreAppTareaEnviada
        from integraciones.moreapp_tareas_ciclo import registrar_ciclo_envio

        defaults = {
            'form_id': 'abc123abc123abc123abc123',
            'form_name': 'Instalacion',
            'task_id': 'task-ciclo-1',
            'mensaje': 'Visita',
            'destinatarios': ['campo@more.app'],
            'creado_por': self.admin,
            'cuenta': '1001',
        }
        defaults.update(kwargs)
        tarea = MoreAppTareaEnviada.objects.create(**defaults)
        registrar_ciclo_envio(tarea)
        return tarea

    def test_informe_con_cuenta_cierra_una_tarea(self):
        from ordenes_trabajo.models import IntegracionMoreApp
        from integraciones.models import MoreAppTareaAlerta, MoreAppTareaEnviada
        from integraciones.moreapp_tareas_ciclo import intentar_completar_por_informe

        tarea = self._tarea()
        otra = self._tarea(task_id='task-ciclo-2', cuenta='2002', mensaje='Otra')
        registro = IntegracionMoreApp.objects.create(
            moreapp_submission_id='sub-ciclo-1',
            nombre_formulario='Instalacion',
            estado_sincronizacion='PROCESADO',
            datos_recibidos={'id': 'sub-ciclo-1', 'info': {'formId': 'abc123abc123abc123abc123', 'formName': 'Instalacion'}},
            datos_procesados={'cliente_codigo': '1001', 'form_name': 'Instalacion'},
        )
        cerrada = intentar_completar_por_informe(registro)
        self.assertEqual(cerrada.pk, tarea.pk)
        tarea.refresh_from_db()
        otra.refresh_from_db()
        self.assertEqual(tarea.estado, MoreAppTareaEnviada.ESTADO_COMPLETADA)
        self.assertEqual(otra.estado, MoreAppTareaEnviada.ESTADO_ASIGNADA)
        self.assertEqual(tarea.integracion_id, registro.pk)
        self.assertTrue(tarea.eventos.filter(tipo='COMPLETADA').exists())
        self.assertTrue(tarea.alertas.filter(tipo=MoreAppTareaAlerta.TIPO_COMPLETADO, broadcast=True).exists())
        self.assertIsNone(intentar_completar_por_informe(registro))

    def test_informe_sin_match_no_altera_tareas(self):
        from ordenes_trabajo.models import IntegracionMoreApp
        from integraciones.models import MoreAppTareaEnviada
        from integraciones.moreapp_tareas_ciclo import intentar_completar_por_informe

        tarea = self._tarea()
        registro = IntegracionMoreApp.objects.create(
            moreapp_submission_id='sub-ciclo-2',
            nombre_formulario='Otro form',
            estado_sincronizacion='PROCESADO',
            datos_recibidos={'info': {'formId': 'zzzzzzzzzzzzzzzzzzzzzzzz', 'formName': 'Otro form'}},
            datos_procesados={'cliente_codigo': '9999', 'form_name': 'Otro form'},
        )
        self.assertIsNone(intentar_completar_por_informe(registro))
        tarea.refresh_from_db()
        self.assertEqual(tarea.estado, MoreAppTareaEnviada.ESTADO_ASIGNADA)

    def test_sync_api_marca_completada(self):
        from integraciones.models import MoreAppTareaEnviada
        from integraciones.moreapp_tareas_ciclo import sincronizar_tareas_desde_api

        tarea = self._tarea()
        with mock.patch(
            'integraciones.moreapp_api.list_form_tasks',
            return_value=[{'id': 'task-ciclo-1', 'status': 'COMPLETED'}],
        ):
            stats = sincronizar_tareas_desde_api()
        tarea.refresh_from_db()
        self.assertEqual(stats['completadas'], 1)
        self.assertEqual(tarea.estado, MoreAppTareaEnviada.ESTADO_COMPLETADA)

    def test_navbar_creador_y_broadcast(self):
        from web.moreapp_avisos import construir_aviso_moreapp
        from django.test import RequestFactory
        from django.contrib.sessions.middleware import SessionMiddleware

        self._tarea(creado_por=self.administrativo)
        factory = RequestFactory()

        def _aviso(user):
            request = factory.get('/')
            request.user = user
            SessionMiddleware(lambda req: None).process_request(request)
            request.session.save()
            return construir_aviso_moreapp(request)

        aviso_creador = _aviso(self.administrativo)
        aviso_admin = _aviso(self.admin)
        aviso_tec = _aviso(self.tecnico)
        self.assertGreaterEqual(aviso_creador['tareas_creacion'], 1)
        self.assertGreaterEqual(aviso_creador['tareas_alertas'], 1)
        self.assertGreaterEqual(aviso_admin['tareas_creacion'], 1)
        self.assertGreaterEqual(aviso_admin['tareas_alertas'], 1)
        self.assertEqual(aviso_tec['tareas_alertas'], 0)

    def test_abrir_listado_marca_alertas_leidas(self):
        from integraciones.models import MoreAppTareaAlerta

        self._tarea()
        self.assertTrue(MoreAppTareaAlerta.objects.filter(leida_en__isnull=True).exists())
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_tareas_list'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Asignada')
        self.assertFalse(MoreAppTareaAlerta.objects.filter(leida_en__isnull=True).exists())

    def test_listado_tiene_formato_terminados_y_acciones(self):
        tarea = self._tarea()
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_tareas_list'))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode('utf-8')
        self.assertIn('Listado', body)
        self.assertIn('Filtrar', body)
        self.assertIn(reverse('moreapp_tarea_detalle', args=[tarea.pk]), body)
        self.assertIn('> Ver</a>', body)

    def test_detalle_muestra_contenido_rellenado_y_pdf(self):
        from ordenes_trabajo.models import IntegracionMoreApp
        from integraciones.models import MoreAppTareaEvento
        from integraciones.moreapp_tareas_ciclo import completar_tarea

        tarea = self._tarea()
        registro = IntegracionMoreApp.objects.create(
            moreapp_submission_id='sub-detalle-1',
            nombre_formulario='Instalacion',
            estado_sincronizacion='PROCESADO',
            datos_recibidos={'data': {'lectura': '1234'}},
            datos_procesados={'cliente_codigo': '1001', 'actividad': 'Cambio medidor'},
        )
        completar_tarea(tarea, origen=MoreAppTareaEvento.ORIGEN_INFORME, integracion=registro)
        self.client.login(username=self.admin.rut, password=self.password)
        detalle = self.client.get(reverse('moreapp_tarea_detalle', args=[tarea.pk]))
        self.assertEqual(detalle.status_code, 200)
        body = detalle.content.decode('utf-8')
        self.assertIn('Contenido rellenado', body)
        self.assertIn('1001', body)
        self.assertIn('Cambio medidor', body)
        self.assertIn(reverse('moreapp_tarea_pdf', args=[tarea.pk]), body)
        pdf = self.client.get(reverse('moreapp_tarea_pdf', args=[tarea.pk]))
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf['Content-Type'], 'application/pdf')
        self.assertTrue(pdf.content.startswith(b'%PDF'))

    def test_pdf_solo_si_completada(self):
        tarea = self._tarea()
        self.client.login(username=self.admin.rut, password=self.password)
        resp = self.client.get(reverse('moreapp_tarea_pdf', args=[tarea.pk]))
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp.url, reverse('moreapp_tarea_detalle', args=[tarea.pk]))
